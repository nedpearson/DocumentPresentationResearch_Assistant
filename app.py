import os
import json
import uuid
import re
import logging
from pathlib import Path
from datetime import datetime

from flask import Flask, render_template, request, jsonify, send_from_directory
from flask_cors import CORS
from flask_sqlalchemy import SQLAlchemy
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from werkzeug.utils import secure_filename
from dotenv import load_dotenv

import anthropic
import PyPDF2
import docx
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

load_dotenv()
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# App & extensions
# ---------------------------------------------------------------------------

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "dev-secret-key-change-in-production")
app.config["SQLALCHEMY_DATABASE_URI"] = os.getenv("DATABASE_URL", "sqlite:///dociq.db")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["MAX_CONTENT_LENGTH"] = int(os.getenv("MAX_CONTENT_LENGTH", 16 * 1024 * 1024))

UPLOAD_FOLDER = os.getenv("UPLOAD_FOLDER", "uploads")
OUTPUT_FOLDER = os.getenv("OUTPUT_FOLDER", "outputs")
ALLOWED_EXTENSIONS = {"pdf", "docx", "doc", "txt", "pptx", "md"}

Path(UPLOAD_FOLDER).mkdir(exist_ok=True)
Path(OUTPUT_FOLDER).mkdir(exist_ok=True)

CORS(app)
db = SQLAlchemy(app)
limiter = Limiter(get_remote_address, app=app, default_limits=[], storage_uri="memory://")

# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------

class Document(db.Model):
    __tablename__ = "documents"
    id            = db.Column(db.String(8),   primary_key=True)
    original_name = db.Column(db.String(255), nullable=False)
    stored_name   = db.Column(db.String(300), nullable=False)
    filepath      = db.Column(db.String(500), nullable=False)
    ext           = db.Column(db.String(10),  nullable=False)
    size          = db.Column(db.Integer,     nullable=False)
    word_count    = db.Column(db.Integer,     default=0)
    char_count    = db.Column(db.Integer,     default=0)
    has_text      = db.Column(db.Boolean,     default=True)
    uploaded_at   = db.Column(db.DateTime,    default=datetime.utcnow)
    text          = db.Column(db.Text,        nullable=True)
    analyses      = db.relationship("AnalysisResult", backref="document", cascade="all, delete-orphan")
    messages      = db.relationship("ChatMessage",    backref="document", cascade="all, delete-orphan")

    def to_dict(self, include_text=False):
        d = {
            "id":             self.id,
            "original_name":  self.original_name,
            "ext":            self.ext,
            "size":           self.size,
            "word_count":     self.word_count,
            "char_count":     self.char_count,
            "has_text":       self.has_text,
            "uploaded_at":    self.uploaded_at.isoformat(),
            "analysis":       any(a.analysis_type == "full" for a in self.analyses),
            "analysis_types": [a.analysis_type for a in self.analyses],
        }
        if include_text:
            d["text"] = self.text or ""
        return d


class AnalysisResult(db.Model):
    __tablename__ = "analysis_results"
    id            = db.Column(db.Integer,  primary_key=True, autoincrement=True)
    doc_id        = db.Column(db.String(8), db.ForeignKey("documents.id"), nullable=False)
    analysis_type = db.Column(db.String(20), nullable=False)
    result        = db.Column(db.Text,    nullable=False)
    created_at    = db.Column(db.DateTime, default=datetime.utcnow)
    model_used    = db.Column(db.String(50), nullable=True)
    __table_args__ = (db.UniqueConstraint("doc_id", "analysis_type", name="uq_doc_atype"),)


class ChatMessage(db.Model):
    __tablename__ = "chat_messages"
    id         = db.Column(db.Integer,    primary_key=True, autoincrement=True)
    session_id = db.Column(db.String(36), nullable=False, index=True)
    doc_id     = db.Column(db.String(8),  db.ForeignKey("documents.id"), nullable=True)
    role       = db.Column(db.String(10), nullable=False)
    content    = db.Column(db.Text,       nullable=False)
    created_at = db.Column(db.DateTime,   default=datetime.utcnow)


with app.app_context():
    db.create_all()

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def get_client():
    key = os.getenv("ANTHROPIC_API_KEY")
    if not key:
        raise ValueError("ANTHROPIC_API_KEY not set. Add it to your .env file.")
    return anthropic.Anthropic(api_key=key)


def extract_text(filepath):
    ext = filepath.rsplit(".", 1)[-1].lower()
    text = ""
    if ext == "pdf":
        with open(filepath, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += (page.extract_text() or "") + "\n"
    elif ext in ("docx", "doc"):
        d = docx.Document(filepath)
        text = "\n".join(p.text for p in d.paragraphs)
    elif ext == "pptx":
        prs = Presentation(filepath)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
    elif ext in ("txt", "md"):
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
    return text.strip()


def trunc(text, n=50000):
    if len(text) <= n:
        return text
    cut = text[:n].rfind(". ")
    cut = cut if cut > int(n * 0.8) else n
    return text[:cut] + "\n\n[... document truncated ...]"


def call_claude(system, user, model="claude-opus-4-6"):
    client = get_client()
    resp = client.messages.create(
        model=model, max_tokens=4096, system=system,
        messages=[{"role": "user", "content": user}],
    )
    return resp.content[0].text


def build_pptx(title, slides_data, output_path):
    DARK = RGBColor(0x0D, 0x1B, 0x2A); ACCENT = RGBColor(0x00, 0xB4, 0xD8)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF); LGRAY = RGBColor(0xCC, 0xCC, 0xCC)
    prs = Presentation()
    prs.slide_width = Inches(13.33); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    sl = prs.slides.add_slide(blank)
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb = DARK
    tb = sl.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(1.5))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run(); run.text = title
    run.font.size = Pt(40); run.font.bold = True; run.font.color.rgb = WHITE
    bar = sl.shapes.add_shape(1, Inches(4.5), Inches(4.2), Inches(4.33), Emu(36000))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
    tb2 = sl.shapes.add_textbox(Inches(1), Inches(4.6), Inches(11.33), Inches(0.5))
    p2 = tb2.text_frame.paragraphs[0]; p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = datetime.now().strftime("%B %d, %Y")
    r2.font.size = Pt(16); r2.font.color.rgb = LGRAY

    for info in slides_data:
        sl = prs.slides.add_slide(blank)
        sl.background.fill.solid(); sl.background.fill.fore_color.rgb = DARK
        hdr = sl.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.33), Emu(180000))
        hdr.fill.solid(); hdr.fill.fore_color.rgb = ACCENT; hdr.line.fill.background()
        tbt = sl.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.33), Inches(0.9))
        pt = tbt.text_frame.paragraphs[0]; rt = pt.add_run()
        rt.text = info.get("title", ""); rt.font.size = Pt(26)
        rt.font.bold = True; rt.font.color.rgb = DARK
        bullets = info.get("content", [])
        if isinstance(bullets, str):
            bullets = [bullets]
        tbc = sl.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(12.33), Inches(5.5))
        tfc = tbc.text_frame; tfc.word_wrap = True
        for i, b in enumerate(bullets):
            para = tfc.paragraphs[i] if i == 0 else tfc.add_paragraph()
            para.space_before = Pt(6); rc = para.add_run()
            rc.text = f"• {b}" if not str(b).startswith("•") else str(b)
            rc.font.size = Pt(18); rc.font.color.rgb = WHITE
    prs.save(output_path)

# ---------------------------------------------------------------------------
# Page routes
# ---------------------------------------------------------------------------

def _stats():
    return {
        "total_docs":    Document.query.count(),
        "analyzed":      db.session.query(db.func.count(db.func.distinct(AnalysisResult.doc_id))).scalar() or 0,
        "total_words":   db.session.query(db.func.sum(Document.word_count)).scalar() or 0,
        "presentations": len(list(Path(OUTPUT_FOLDER).glob("*.pptx"))),
        "chats":         db.session.query(db.func.count(db.func.distinct(ChatMessage.session_id))).scalar() or 0,
    }


@app.route("/")
def index():
    recent = [d.to_dict() for d in Document.query.order_by(Document.uploaded_at.desc()).limit(5).all()]
    return render_template("index.html", stats=_stats(), recent=recent)


@app.route("/analyze")
def analyze_page():
    return render_template("analyze.html", documents=[d.to_dict() for d in Document.query.order_by(Document.uploaded_at.desc()).all()])


@app.route("/research")
def research_page():
    return render_template("research.html", documents=[d.to_dict() for d in Document.query.order_by(Document.uploaded_at.desc()).all()])


@app.route("/present")
def present_page():
    pptx_files = [f.name for f in Path(OUTPUT_FOLDER).glob("*.pptx")]
    return render_template("present.html",
        documents=[d.to_dict() for d in Document.query.order_by(Document.uploaded_at.desc()).all()],
        presentations=pptx_files)


@app.route("/documents")
def documents_page():
    return render_template("documents.html", documents=[d.to_dict() for d in Document.query.order_by(Document.uploaded_at.desc()).all()])

# ---------------------------------------------------------------------------
# API – Document management
# ---------------------------------------------------------------------------

@app.route("/api/upload", methods=["POST"])
def upload():
    if "file" not in request.files:
        return jsonify({"error": "No file provided"}), 400
    file = request.files["file"]
    if not file.filename or not allowed_file(file.filename):
        return jsonify({"error": "Unsupported file type. Allowed: pdf, docx, txt, pptx, md"}), 400

    doc_id   = str(uuid.uuid4())[:8]
    filename = secure_filename(file.filename)
    stored   = f"{doc_id}_{filename}"
    filepath = os.path.join(UPLOAD_FOLDER, stored)
    file.save(filepath)

    warning = None
    try:
        text = extract_text(filepath)
    except Exception as e:
        logger.error(f"Text extraction failed for {filename}: {e}")
        text = ""
        warning = "Text could not be extracted from this file (may be image-based or corrupted)."

    if not text and not warning:
        warning = "No extractable text found. The file may be image-based or empty."

    doc = Document(
        id=doc_id, original_name=filename, stored_name=stored, filepath=filepath,
        ext=filename.rsplit(".", 1)[-1].lower(), size=os.path.getsize(filepath),
        word_count=len(text.split()) if text else 0, char_count=len(text),
        has_text=bool(text), text=text,
    )
    db.session.add(doc)
    db.session.commit()

    resp = {"id": doc_id, "name": filename, "word_count": doc.word_count, "size": doc.size, "has_text": doc.has_text}
    if warning:
        resp["warning"] = warning
    return jsonify(resp)


@app.route("/api/documents")
def list_docs():
    return jsonify({"documents": [d.to_dict() for d in Document.query.order_by(Document.uploaded_at.desc()).all()]})


@app.route("/api/documents/<doc_id>", methods=["DELETE"])
def delete_doc(doc_id):
    doc = Document.query.get(doc_id)
    if not doc:
        return jsonify({"error": "Document not found"}), 404
    filepath = doc.filepath
    db.session.delete(doc)
    db.session.commit()
    try:
        os.remove(filepath)
    except OSError:
        pass
    return jsonify({"message": "Document deleted"})


@app.route("/api/documents/<doc_id>/text")
def get_doc_text(doc_id):
    """Return raw extracted text for the preview modal."""
    doc = Document.query.get(doc_id)
    if not doc:
        return jsonify({"error": "Document not found"}), 404
    return jsonify({
        "id": doc.id, "name": doc.original_name,
        "text": doc.text or "", "word_count": doc.word_count,
        "char_count": doc.char_count, "has_text": doc.has_text,
    })

# ---------------------------------------------------------------------------
# API – Analysis (with per-type caching)
# ---------------------------------------------------------------------------

ANALYSIS_PROMPTS = {
    "full": (
        "Perform a comprehensive analysis.\n"
        "## Executive Summary\n## Key Themes\n## Main Arguments\n"
        "## Evidence & Data\n## Strengths & Gaps\n## Key Takeaways\n## Recommended Actions"
    ),
    "summary":   "Write a concise 3-5 paragraph executive summary covering purpose, main points, and conclusions.",
    "insights":  "Extract the top 10 key insights as numbered items with bold titles and brief explanations.",
    "sentiment": ("Analyze tone, sentiment, and rhetorical style: overall tone, emotional language, "
                  "persuasion techniques, objectivity level, and intended audience."),
    "entities":  ("Extract and categorize all key entities:\n"
                  "**People** | **Organizations** | **Locations** | **Dates/Times** | **Key Terms** | **Statistics**"),
}


@app.route("/api/analyze/<doc_id>", methods=["POST"])
@limiter.limit("20 per hour")
def analyze(doc_id):
    doc = Document.query.get(doc_id)
    if not doc:
        return jsonify({"error": "Document not found"}), 404
    if not doc.has_text or not doc.text:
        return jsonify({"error": "This document has no extractable text (may be image-based)."}), 400

    data  = request.json or {}
    atype = data.get("type", "full")
    force = data.get("force", False)

    if atype not in ANALYSIS_PROMPTS:
        return jsonify({"error": f"Unknown analysis type '{atype}'"}), 400

    # Cache check — skip if force=true
    if not force:
        cached = AnalysisResult.query.filter_by(doc_id=doc_id, analysis_type=atype).first()
        if cached:
            logger.info(f"Cache hit: doc={doc_id} type={atype}")
            return jsonify({"result": cached.result, "type": atype, "cached": True})

    prompt = f"{ANALYSIS_PROMPTS[atype]}\n\nDocument content:\n\n{trunc(doc.text)}"
    system = ("You are an expert document analyst. Provide structured, insightful analysis. "
              "Use markdown headers, bullet points, and bold text for clarity.")
    try:
        result = call_claude(system, prompt)

        # Upsert
        existing = AnalysisResult.query.filter_by(doc_id=doc_id, analysis_type=atype).first()
        if existing:
            existing.result = result; existing.created_at = datetime.utcnow()
        else:
            db.session.add(AnalysisResult(doc_id=doc_id, analysis_type=atype,
                                          result=result, model_used="claude-opus-4-6"))
        db.session.commit()
        return jsonify({"result": result, "type": atype, "cached": False})
    except Exception as e:
        logger.error(f"Analysis error doc={doc_id} type={atype}: {e}")
        return jsonify({"error": str(e)}), 500

# ---------------------------------------------------------------------------
# API – Research chat (persisted)
# ---------------------------------------------------------------------------

@app.route("/api/chat", methods=["POST"])
@limiter.limit("30 per hour")
def chat():
    data       = request.json or {}
    doc_ids    = data.get("doc_ids", [])
    question   = data.get("question", "").strip()
    session_id = data.get("session_id") or str(uuid.uuid4())

    if not question:
        return jsonify({"error": "Question is required"}), 400
    if len(question) > 4000:
        return jsonify({"error": "Question too long (max 4,000 characters)"}), 400

    ctx_parts    = []
    primary_did  = None
    for did in doc_ids:
        doc = Document.query.get(did)
        if doc and doc.text:
            ctx_parts.append(f"### {doc.original_name}\n\n{trunc(doc.text, 20000)}")
            if primary_did is None:
                primary_did = did
    context = "\n\n---\n\n".join(ctx_parts) or "No documents selected."

    system = (
        "You are an expert research assistant. You have access to the following document(s):\n\n"
        f"{context}\n\n"
        "Answer questions accurately based on the content. Cite specific sections when relevant. "
        "If information is not in the documents, say so clearly. Use markdown for clarity."
    )
    try:
        answer = call_claude(system, question)
        db.session.add(ChatMessage(session_id=session_id, role="user",      content=question, doc_id=primary_did))
        db.session.add(ChatMessage(session_id=session_id, role="assistant", content=answer,   doc_id=primary_did))
        db.session.commit()
        return jsonify({"answer": answer, "session_id": session_id})
    except Exception as e:
        logger.error(f"Chat error session={session_id}: {e}")
        return jsonify({"error": str(e)}), 500

# ---------------------------------------------------------------------------
# API – Presentation generator
# ---------------------------------------------------------------------------

@app.route("/api/generate-presentation", methods=["POST"])
@limiter.limit("5 per hour")
def gen_presentation():
    data       = request.json or {}
    doc_ids    = data.get("doc_ids", [])
    title      = (data.get("title") or "Presentation").strip()
    num_slides = max(3, min(int(data.get("num_slides", 8)), 20))
    style      = data.get("style", "professional")

    if not doc_ids:
        return jsonify({"error": "Select at least one document"}), 400

    combined = ""
    for did in doc_ids:
        doc = Document.query.get(did)
        if doc and doc.text:
            combined += f"\n\n=== {doc.original_name} ===\n\n{doc.text}"

    if not combined.strip():
        return jsonify({"error": "Selected documents have no extractable text"}), 400

    system = ("You are a professional presentation designer. "
              "Create compelling, well-structured slides that communicate key ideas clearly.")
    prompt = (
        f"Create a {style} presentation titled '{title}' with exactly {num_slides} content slides "
        f"based on:\n\n{trunc(combined, 40000)}\n\n"
        "Return ONLY valid JSON (no markdown fences, no extra text):\n"
        '{"slides":[{"title":"...","content":["bullet1","bullet2","bullet3"]}]}\n'
        "Each slide: 3-5 concise bullet points. Titles must be compelling and actionable."
    )
    try:
        raw = call_claude(system, prompt)
        raw = re.sub(r"```(?:json)?\s*", "", raw).strip()
        m = re.search(r'\{[\s\S]*\}', raw)
        if not m:
            raise ValueError("Could not parse slide structure from AI response")
        slides_data = json.loads(m.group())["slides"]

        pid   = str(uuid.uuid4())[:8]
        safe  = secure_filename(title)[:40] or "presentation"
        fname = f"{pid}_{safe}.pptx"
        build_pptx(title, slides_data, os.path.join(OUTPUT_FOLDER, fname))

        return jsonify({"message": "Presentation generated", "filename": fname,
                        "slides": len(slides_data), "download_url": f"/api/download/{fname}"})
    except Exception as e:
        logger.error(f"Presentation error: {e}")
        return jsonify({"error": str(e)}), 500


@app.route("/api/download/<filename>")
def download(filename):
    return send_from_directory(OUTPUT_FOLDER, secure_filename(filename), as_attachment=True)


@app.route("/api/stats")
def stats():
    return jsonify(_stats())

# ---------------------------------------------------------------------------
# Error handlers
# ---------------------------------------------------------------------------

@app.errorhandler(413)
def too_large(e):
    return jsonify({"error": "File too large. Maximum size is 16 MB."}), 413

@app.errorhandler(429)
def rate_limited(e):
    return jsonify({"error": "Rate limit reached. Please wait before making more AI requests."}), 429

@app.errorhandler(404)
def not_found(e):
    return render_template("404.html"), 404


if __name__ == "__main__":
    debug = os.getenv("FLASK_DEBUG", "true").lower() == "true"
    port  = int(os.getenv("PORT", 5000))
    app.run(debug=debug, host="0.0.0.0", port=port)
